# -*- coding: utf-8 -*-
"""单用途工具：读一份排班 xlsx，打印「结构 + 班次 + 底色 + 每日覆盖」摘要。

只读，不修改文件，也不做任何排班判断。

用法：
    python 读排班表.py --xlsx "2026年10月客服排班表.xlsx"
    python 读排班表.py --xlsx xxx.xlsx --sheet 2026年10月 --json 摘要.json

组合用法：本工具只负责“读”，把事实交给经验和人判断；写格、刷色、对拍等
操作另用其他单一用途小工具，按需组合，不做一键整月排班。
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

# 结构行/非员工行：这些“姓名”不是员工
SKIP_NAMES = {
    "日期", "星期", "月份", "职务", "售前", "售后", "早班", "晚班", "休息",
    "年假", "行政", "剩余", "实到", "应到", "备注", "上班人数", "本月休", "本月天",
}


def cell_color(cell) -> str:
    """取有效填充色的后 6 位（大写）；无填充返回空串。"""
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    rgb = fill.start_color.rgb
    if not rgb:
        return ""
    text = str(rgb)
    return text[-6:].upper() if len(text) >= 6 else text


def read_schedule(path: str, sheet: str | None = None) -> dict:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.active

    header_row = None
    name_col = None
    day_cols: dict[int, int] = {}
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)), start=1):
        values = [c.value for c in row]
        if any(v is not None and str(v).strip() == "日期" for v in values):
            header_row = r
            name_col = next(i for i, v in enumerate(values, start=1)
                            if v is not None and str(v).strip() == "日期")
            for i, v in enumerate(values, start=1):
                text = str(v).strip() if v is not None else ""
                if text.isdigit() and 1 <= int(text) <= 31:
                    day_cols.setdefault(int(text), i)
            break
    if header_row is None or name_col is None:
        raise SystemExit("没找到含「日期」的表头行")
    days = sorted(day_cols)

    employees = []
    last_group = ""
    for r in range(header_row + 1, ws.max_row + 1):
        group = ws.cell(r, 1).value
        if group is not None and str(group).strip():
            last_group = str(group).strip()
        name = ws.cell(r, name_col).value
        name = str(name).strip() if name is not None else ""
        if not name or name in SKIP_NAMES:
            continue
        if any(key in name for key in ("小计", "汇总", "人数", "本月", "应休")):
            continue
        shifts, colors = {}, {}
        for d in days:
            cell = ws.cell(r, day_cols[d])
            value = cell.value
            shifts[d] = str(value).strip() if value is not None else ""
            color = cell_color(cell)
            if color:
                colors[d] = color
        employees.append({"row": r, "group": last_group, "name": name,
                          "shifts": shifts, "colors": colors})

    coverage = {}
    for group in sorted({e["group"] for e in employees if e["group"]}):
        members = [e for e in employees if e["group"] == group]
        coverage[group] = {
            d: {
                "早": sum(1 for e in members if e["shifts"][d] == "早"),
                "晚": sum(1 for e in members if e["shifts"][d] == "晚"),
                "休": sum(1 for e in members if e["shifts"][d] == ""),
                "其他": sum(1 for e in members if e["shifts"][d] not in ("早", "晚", "")),
            }
            for d in days
        }

    return {
        "file": str(Path(path).resolve()),
        "sheet": ws.title,
        "days": days,
        "employees": employees,
        "coverage": coverage,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="读排班 xlsx 并打印摘要（只读）")
    parser.add_argument("--xlsx", required=True, help="排班表 xlsx 路径")
    parser.add_argument("--sheet", help="工作表名（默认活动表）")
    parser.add_argument("--json", help="可选：把完整摘要写到 JSON")
    args = parser.parse_args()

    data = read_schedule(args.xlsx, args.sheet)
    days = data["days"]

    print(f"文件: {data['file']}")
    print(f"工作表: {data['sheet']}（共 {len(days)} 天）")
    print(f"员工 {len(data['employees'])} 人: " + "、".join(e["name"] for e in data["employees"]))

    for group, per_day in data["coverage"].items():
        for key in ("早", "晚", "休", "其他"):
            counts = [per_day[d][key] for d in days]
            if len(set(counts)) == 1:
                span = str(counts[0])
            else:
                span = f"{min(counts)}~{max(counts)}"
            print(f"每日覆盖[{group}] {key}: {span}")

    print("\n每人摘要:")
    for e in data["employees"]:
        shifts = e["shifts"]
        counts = {code: sum(1 for d in days if shifts[d] == code) for code in ("早", "晚", "年", "行")}
        rest = sum(1 for d in days if shifts[d] == "")
        color_days = ", ".join(f"{d}={c}" for d, c in sorted(e["colors"].items())) or "无"
        print(f"  {e['name']}[{e['group']}] 早{counts['早']} 晚{counts['晚']} 休{rest} "
              f"年{counts['年']} 行{counts['行']} | 底色: {color_days}")

    if args.json:
        with open(args.json, "w", encoding="utf-8") as handle:
            json.dump(data, handle, ensure_ascii=False, indent=2)
        print(f"\n已写出 JSON: {args.json}")


if __name__ == "__main__":
    main()
