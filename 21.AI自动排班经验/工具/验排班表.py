# -*- coding: utf-8 -*-
"""单用途工具：按《验收清单.md》逐项校验一份排班表（只读，不改文件）。

吃两种输入：排班 xlsx（能读底色和公式）和导出/快照 TSV（只有数值，跳过值班检查）。
上月用 --prev 传入（xlsx 或快照 TSV 都行），用于跨月衔接校验。

统计列（休息/早班/晚班/实到/应到）如果是公式，按“公式结构”验：
必须覆盖整月（`$C4:$AG4`）或按当月天数动态取范围（`OFFSET` + 引用「本月天数：」格），
避免 30 天的月份把 31 号空白算成休息；如果是数值，则直接和矩阵对账。
售后值班也要验：每班（早/晚）只标 1 人浅蓝 `BDD7EE`，组长李守耀在岗时他本人标黄 `FFFF00`。

用法：
    python 验排班表.py --xlsx "测试数据\\2026年10月客服排班表.xlsx" --lead 李守耀
    python 验排班表.py --xlsx 本月.xlsx --prev "……\\2026年9月.tsv" --lead 李守耀 --json 报告.json

只负责“判”；怎么改还是人（或别的单一用途小工具）的事。
"""
from __future__ import annotations

import argparse
import csv
import importlib
import json
import math
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
_reader = importlib.import_module("读排班表")
SKIP_NAMES = set(_reader.SKIP_NAMES)

GREEN = "E2F0D9"
BLUE_AFTER = "BDD7EE"    # 售后值班（每班 1 人）
YELLOW_LEAD = "FFFF00"   # 售后组长（李守耀）在岗
WORK = ("早", "晚", "行")
KINDS = ("早班", "晚班", "休息")
POLICY_LABELS = ("本月休息天数：", "本月天数：")
STAT_KEYS = ("剩余", "年假", "休息", "早班", "晚班", "实到", "应到")


# ---------------------------------------------------------------- 读表

class XlsxGrid:
    def __init__(self, path: str, sheet: str | None = None):
        wb = load_workbook(path, data_only=False)  # 取公式文本
        self.ws = wb[sheet] if sheet else wb.active
        self.max_row, self.max_col = self.ws.max_row, self.ws.max_column

    def get(self, r: int, c: int):
        return self.ws.cell(r, c).value

    def color(self, r: int, c: int) -> str:
        fill = self.ws.cell(r, c).fill
        if not fill or fill.fill_type != "solid":
            return ""
        rgb = fill.start_color.rgb
        text = str(rgb) if rgb else ""
        return text[-6:].upper() if len(text) >= 6 else ""


class TsvGrid:
    def __init__(self, path: str):
        with open(path, encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle, delimiter="\t"))
        self.rows = [[(c or "").strip() for c in row] for row in rows]
        self.max_row = len(self.rows)
        self.max_col = max((len(r) for r in self.rows), default=0)

    def get(self, r: int, c: int):
        row = self.rows[r - 1] if 0 < r <= self.max_row else []
        value = row[c - 1] if 0 < c <= len(row) else ""
        return value if value != "" else None

    def color(self, r: int, c: int) -> str:
        return ""


def load_grid(path: str, sheet: str | None = None):
    return TsvGrid(path) if path.lower().endswith((".tsv", ".txt")) else XlsxGrid(path, sheet)


def text(value) -> str:
    return str(value).strip() if value is not None else ""


def is_formula(value) -> bool:
    return isinstance(value, str) and value.startswith("=")


def norm(formula: str) -> str:
    """去空格、去 $，方便比对公式里引用的范围。"""
    return formula.replace(" ", "").replace("$", "").upper()


def as_int(value):
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None


def parse(grid) -> dict:
    header = name_col = None
    day_col: dict[int, int] = {}
    for r in range(1, min(grid.max_row, 20) + 1):
        vals = [text(grid.get(r, c)) for c in range(1, grid.max_col + 1)]
        if "日期" in vals:
            header = r
            name_col = vals.index("日期") + 1
            for i, v in enumerate(vals, start=1):
                if v.isdigit() and 1 <= int(v) <= 31:
                    day_col.setdefault(int(v), i)
            break
    if header is None:
        raise SystemExit("没找到含「日期」的表头行")
    days = sorted(day_col)

    stat_col: dict[str, int] = {}
    policy: dict[str, dict] = {}
    for r in range(1, min(grid.max_row, 8) + 1):
        for c in range(name_col + 1, grid.max_col + 1):
            v = text(grid.get(r, c))
            if v in STAT_KEYS:
                stat_col.setdefault(v, c)
    for r in range(1, grid.max_row + 1):
        for c in range(1, grid.max_col + 1):
            v = text(grid.get(r, c))
            if v in POLICY_LABELS:
                policy[v] = {"row": r, "col": c + 1, "cell": grid.get(r, c + 1)}

    employees, group = [], ""
    for r in range(header + 1, grid.max_row + 1):
        g = text(grid.get(r, 1))
        if g in ("售前", "售后"):
            group = g
        name = text(grid.get(r, name_col))
        if not name or name in SKIP_NAMES or name in KINDS or "小计" in name or "汇总" in name:
            continue
        employees.append({
            "row": r, "name": name, "group": group,
            "shifts": {d: text(grid.get(r, day_col[d])) for d in days},
            "colors": {d: grid.color(r, day_col[d]) for d in days},
            "stats": {k: grid.get(r, c) for k, c in stat_col.items()},
        })

    summary: dict[tuple[str, str], dict[int, object]] = {}
    head_count: dict[int, object] = {}
    subtotal: dict[str, dict[str, object]] = {}
    marks = []
    current = kind = ""
    for r in range(header + 1, grid.max_row + 1):
        g1, g2 = text(grid.get(r, 1)), text(grid.get(r, 2))
        if g1 in ("售前", "售后") and g2 in KINDS:
            current, kind = g1, g2
        elif g2 in KINDS:
            kind = g2
        else:
            kind = ""
        if kind and current:
            summary[(current, kind)] = {d: grid.get(r, day_col[d]) for d in days}
        if g2 == "上班人数":
            head_count = {d: grid.get(r, day_col[d]) for d in days}
        whole = [text(grid.get(r, c)) for c in range(1, grid.max_col + 1)]
        for c, v in enumerate(whole, start=1):
            if v.endswith("小计"):
                subtotal[v[:-2]] = {k: grid.get(r, col) for k, col in stat_col.items()}
            if re.fullmatch(r"\d{1,2}月", v):
                marks.append((r, c, grid.color(r, c)))

    formulas = [s for e in employees for s in e["shifts"].values() if is_formula(s)]

    return {"days": days, "employees": employees, "summary": summary, "head_count": head_count,
            "subtotal": subtotal, "marks": marks, "stat_col": stat_col, "policy": policy,
            "day_formulas": len(formulas)}


# ---------------------------------------------------------------- 校验

class Report:
    def __init__(self):
        self.items: list[tuple[str, str, str]] = []

    def add(self, section, level, msg):
        self.items.append((section, level, msg))

    def errors(self):
        return [i for i in self.items if i[1] == "error"]


def groups_of(employees):
    out: dict[str, list[dict]] = {}
    for e in employees:
        out.setdefault(e["group"], []).append(e)
    return out


def check_cross(cur, prev, report, max_streak):
    days, pdays = cur["days"], prev["days"]
    pmap = {e["name"]: e for e in prev["employees"]}
    last = pdays[-1]
    lines = []
    for e in cur["employees"]:
        pe = pmap.get(e["name"])
        if not pe:
            report.add("〇 跨月衔接", "warn", f"{e['name']} 在上月表里没找到，跨月无法核对")
            continue
        pv = pe["shifts"][last]
        trail = 0
        for d in reversed(pdays):
            if pe["shifts"][d] in WORK:
                trail += 1
            else:
                break
        lead = 0
        for d in days:
            if e["shifts"][d] in WORK:
                lead += 1
            else:
                break
        first = e["shifts"][days[0]]
        if pv == "晚" and first == "早":
            report.add("〇 跨月衔接", "error", f"{e['name']} 跨月晚转早（上月末晚 → 本月 1 号早）")
        total = trail + lead
        if lead == 0:
            pass  # 本月一上班就休，上月遗留的连上已截止，不再报
        elif total > max_streak + 1:
            report.add("〇 跨月衔接", "error",
                       f"{e['name']} 跨月连上 {trail}(上月) + {lead}(本月) = {total} > {max_streak + 1}")
        elif total > max_streak:
            report.add("〇 跨月衔接", "warn",
                       f"{e['name']} 跨月连上 {trail}+{lead} = {total}（上限 {max_streak}，需说明）")
        lines.append(f"{e['name']} 上月{last}日{pv or '休'}(连{trail}) → 本月头{lead}")
    report.add("〇 跨月衔接", "info", "；".join(lines))


def check_coverage(cur, report, lead):
    for group, members in groups_of(cur["employees"]).items():
        for d in cur["days"]:
            cnt = {code: sum(1 for e in members if e["shifts"][d] == code) for code in ("早", "晚")}
            rest = sum(1 for e in members if e["shifts"][d] == "")
            if group == "售前":
                if cnt["早"] != 2 or cnt["晚"] != 2 or rest != 1:
                    report.add("一 在岗覆盖", "error",
                               f"售前 d{d}: 早{cnt['早']} 晚{cnt['晚']} 休{rest}（应为 2/2/1）")
            else:
                if cnt["早"] < 2:
                    report.add("一 在岗覆盖", "error", f"售后 d{d}: 早{cnt['早']}（保底 2 早）")
                if not 1 <= cnt["晚"] <= 2:
                    report.add("一 在岗覆盖", "warn", f"售后 d{d}: 晚{cnt['晚']}（一般 1~2）")
                others = [e for e in members if e["name"] != lead]
                if lead and any(e["shifts"][d] == "" for e in members if e["name"] == lead):
                    idle = [e["name"] for e in others if e["shifts"][d] == ""]
                    if idle:
                        report.add("一 在岗覆盖", "error",
                                   f"售后 d{d} {lead}休，但 {('、'.join(idle))} 也休（应 4 人全上）")
    if lead:
        work = [e for e in cur["employees"] if e["name"] == lead]
        if work and any(e["shifts"][d] == "晚" for e in work for d in cur["days"]):
            report.add("一 在岗覆盖", "error", f"{lead} 出现晚班（应全早）")


def check_shifts(cur, report, max_streak):
    for e in cur["employees"]:
        days = cur["days"]
        run = block = 0
        block_start = days[0]
        for i, d in enumerate(days):
            v = e["shifts"][d]
            run = run + 1 if v in WORK else 0
            if run > max_streak + 1:
                report.add("三 连上与休息", "error", f"{e['name']} d{d} 连上 {run} 天（上限 {max_streak}）")
            elif run > max_streak:
                report.add("三 连上与休息", "warn", f"{e['name']} d{d} 连上 {run} 天（需说明）")
            if v == "晚":
                block += 1
            else:
                if block > 4:
                    report.add("二 晚班", "warn",
                               f"{e['name']} d{block_start}~d{days[i - 1]} 连续晚班 {block} 天（一般 ≤4）")
                block = 0
                block_start = d
            if i and e["shifts"][days[i - 1]] == "晚" and v == "早":
                report.add("二 晚班", "error", f"{e['name']} d{d} 晚接早")
        if block > 4:
            report.add("二 晚班", "warn", f"{e['name']} d{block_start}~d{days[-1]} 连续晚班 {block} 天")


def check_duty(cur, report, green):
    if all(not c for e in cur["employees"] for c in e["colors"].values()):
        report.add("四 值班", "warn", "输入没有底色数据（TSV），值班检查跳过")
        return
    sellers = [e for e in cur["employees"] if e["group"] == "售前"]
    if not sellers:
        return
    duty = {d: [e["name"] for e in sellers if e["colors"][d] == green] for d in cur["days"]}
    bad_days = [d for d in cur["days"] if len(duty[d]) != 2]
    for d in bad_days:
        report.add("四 值班", "error", f"d{d} 绿标 {len(duty[d])} 个（每天早、晚各 1）：{duty[d]}")
    for d in cur["days"]:
        if len(duty[d]) == 2:
            shifts = sorted(e["shifts"][d] for e in sellers if e["name"] in duty[d])
            if shifts != ["早", "晚"]:
                report.add("四 值班", "error", f"d{d} 值班班次 {shifts}（应一早在岗、一晚在岗）")
    if bad_days:
        return
    total, n = len(cur["days"]) * 2, len(sellers)
    lo, hi = math.floor(total / n), math.ceil(total / n)
    for e in sellers:
        ds = [d for d in cur["days"] if e["name"] in duty[d]]
        for a, b in zip(ds, ds[1:]):
            if b == a + 1:
                report.add("四 值班", "error", f"{e['name']} d{a}→d{b} 连续两天值班")
        if not lo - 1 <= len(ds) <= hi + 1:
            report.add("四 值班", "error", f"{e['name']} 值班 {len(ds)} 次（应 {lo}~{hi}）")
        elif not lo <= len(ds) <= hi:
            report.add("四 值班", "warn", f"{e['name']} 值班 {len(ds)} 次（一般 {lo}~{hi}）")


def check_marks(cur, report):
    for e in cur["employees"]:
        if any(e["shifts"][d] == "年" for d in cur["days"]):
            report.add("五 特殊标记", "error", f"{e['name']} 出现「年」（年假不主动排）")
        for d in cur["days"]:
            if e["shifts"][d] in ("行", "年") and e["colors"][d] == GREEN:
                report.add("五 特殊标记", "error", f"{e['name']} d{d} 休息却带值班绿")
    if cur["day_formulas"]:
        report.add("五 特殊标记", "warn", f"有 {cur['day_formulas']} 个班次格是公式，可能读不出班次")
    if not cur["marks"]:
        report.add("五 特殊标记", "warn", "没找到「X月」月份标记")
    for r, c, color in cur["marks"]:
        if color and color != "FF0000":
            report.add("五 特殊标记", "warn", f"月份标记（行{r}列{c}）颜色 {color} 不是红")


def check_after_duty(cur, report, lead):
    """售后值班：每班（早/晚）只标 1 人浅蓝；组长在岗时他自己是值班负责人（黄）。"""
    if all(not c for e in cur["employees"] for c in e["colors"].values()):
        return   # TSV：check_duty 已提醒
    after = [e for e in cur["employees"] if e["group"] == "售后"]
    if not after:
        return
    multi_e, multi_l, miss_e, miss_l, off_shift, lead_blue, lead_gray = [], [], [], [], [], [], []
    for d in cur["days"]:
        early = [e for e in after if e["shifts"][d] == "早"]
        late = [e for e in after if e["shifts"][d] == "晚"]
        blue_e = [e["name"] for e in early if e["colors"][d] == BLUE_AFTER]
        blue_l = [e["name"] for e in late if e["colors"][d] == BLUE_AFTER]
        lead_e = next((e for e in after if e["name"] == lead), None)
        lead_on = bool(lead_e and lead_e["shifts"][d] in ("早", "晚"))
        if len(blue_e) > 1:
            multi_e.append(f"d{d} {len(blue_e)}人")
        if len(blue_l) > 1:
            multi_l.append(f"d{d} {len(blue_l)}人")
        if early and not lead_on and not blue_e:
            miss_e.append(f"d{d}")
        if late and not blue_l:
            miss_l.append(f"d{d}")
        if lead_on and blue_e:
            lead_blue.append(f"d{d} {blue_e}")
        for e in after:
            if e["colors"][d] == BLUE_AFTER and e["shifts"][d] not in ("早", "晚"):
                off_shift.append(f"{e['name']} d{d}({e['shifts'][d] or '休'})")
            if lead_e and e["name"] == lead and e["shifts"][d] in ("早", "晚") and e["colors"][d] != YELLOW_LEAD:
                lead_gray.append(f"d{d}")

    def brief(items, limit=8):
        return "、".join(items[:limit]) + (f" …共{len(items)}天" if len(items) > limit else "")

    if multi_e:
        report.add("四 值班", "error", f"售后早班浅蓝超过 1 人（每班只标 1 人）：{brief(multi_e)}")
    if multi_l:
        report.add("四 值班", "error", f"售后晚班浅蓝超过 1 人（每班只标 1 人）：{brief(multi_l)}")
    if miss_e:
        report.add("四 值班", "error", f"售后早班没标值班（组长不在的日子应从早班选 1 人）：{brief(miss_e)}")
    if miss_l:
        report.add("四 值班", "error", f"售后晚班没标值班（应从晚班选 1 人）：{brief(miss_l)}")
    if off_shift:
        report.add("四 值班", "error", f"售后值班标在非在岗格上：{brief(off_shift)}")
    if lead_blue:
        report.add("四 值班", "warn", f"组长在班、早班又标了浅蓝：{brief(lead_blue)}")
    if lead_gray:
        report.add("四 值班", "warn", f"{lead} 在岗但没标组长色 {YELLOW_LEAD}：{brief(lead_gray)}")

    counts = {e["name"]: sum(1 for d in cur["days"] if e["colors"][d] == BLUE_AFTER) for e in after}
    report.add("四 值班", "info", "售后值班（浅蓝）次数：" + "、".join(f"{k} {v} 次" for k, v in counts.items() if v))
    for e in after:
        if multi_e or multi_l:      # 整片涂错时，连值提醒只会刷屏
            break
        ds = [d for d in cur["days"] if e["colors"][d] == BLUE_AFTER]
        for a, b in zip(ds, ds[1:]):
            if b != a + 1:
                continue
            same_shift = [x for x in after if x["shifts"][b] == e["shifts"][b]]
            if len(same_shift) > 1:      # 那天同班次还有别人可换，才算是“没换人”
                report.add("四 值班", "warn", f"{e['name']} d{a}→d{b} 连着两天售后值班（能间隔就间隔换人）")


def check_stats(cur, report):
    if not cur["stat_col"]:
        report.add("六 数据统计", "warn", "没找到统计列（剩余/年假/休息/早班/晚班/实到/应到）")
        return
    days = cur["days"]
    n = len(days)
    policy = cur["policy"]
    day_cell = policy.get("本月天数：")
    rest_cell = policy.get("本月休息天数：")

    def ref(cellinfo):
        return f"{get_column_letter(cellinfo['col'])}{cellinfo['row']}" if cellinfo else None

    day_ref, rest_ref = ref(day_cell), ref(rest_cell)
    # 「本月天数」必须自动算当月天数，否则 30 天的月份会把 31 号空白算成休息
    if day_cell:
        got = day_cell["cell"]
        if is_formula(got):
            f = norm(got)
            if "COUNTA(" not in f and "COUNT(" not in f:
                report.add("六 数据统计", "error", f"「本月天数」公式没数天数：{got}")
            if "C2" not in f or "AG2" not in f:
                report.add("六 数据统计", "warn", f"「本月天数」公式没引用日期行 C2:AG2：{got}")
        elif as_int(got) != n:
            report.add("六 数据统计", "warn",
                       f"「本月天数」写死 {got}，当月实际 {n} 天，统计公式可能把多的那列算成休息")
    else:
        report.add("六 数据统计", "warn", "没找到「本月天数：」格（统计公式是否按当月天数取范围要人工看）")

    want_keys = ("年假", "休息", "早班", "晚班", "实到", "应到")
    formula_ok = value_ok = 0
    for e in cur["employees"]:
        want = {
            "年假": sum(1 for d in days if e["shifts"][d] == "年"),
            "休息": sum(1 for d in days if e["shifts"][d] == ""),
            "早班": sum(1 for d in days if e["shifts"][d] == "早"),
            "晚班": sum(1 for d in days if e["shifts"][d] == "晚"),
            "实到": sum(1 for d in days if e["shifts"][d] in WORK),
            "应到": n - (as_int(rest_cell["cell"]) if rest_cell and not is_formula(rest_cell["cell"]) else 0),
        }
        for key in want_keys:
            got = e["stats"].get(key)
            if got is None:
                continue
            if is_formula(got):
                f = norm(got)
                bad = []
                row = e["row"]
                month_range = f"C{row}:AG{row}"          # 明确覆盖整月（1号..31号）
                day_ref_n = norm(f"={day_ref}").strip("=") if day_ref else ""
                has_month_range = month_range in f
                has_dynamic = "OFFSET(" in f and day_ref_n and day_ref_n in f
                def stat_ref(key):
                    col = cur["stat_col"].get(key)
                    return f"{get_column_letter(col)}{row}" if col else ""
                if key == "应到":
                    if rest_ref and norm(f"={rest_ref}").strip("=") not in f:
                        bad.append(f"没引用「本月休息天数」{rest_ref}")
                    if day_ref_n and day_ref_n not in f:
                        bad.append(f"没引用「本月天数」{day_ref}")
                else:
                    if not (has_month_range or has_dynamic):
                        bad.append(f"没覆盖整月（既不是 $C{row}:$AG{row}，也没按「本月天数」动态取范围）")
                    if key == "休息":
                        # 必须引用「本月天数」：无论 COUNTBLANK(OFFSET(...)) 还是 =天数−早−晚−年−行；
                        # 只写 COUNTBLANK($C4:$AG4) 在 30 天的月份会把 31 号空白算成休息
                        if not day_ref_n or day_ref_n not in f:
                            bad.append("休息没引用「本月天数」（30 天的月份会把空白列算成休息）")
                        elif "COUNTBLANK(" not in f and "-" not in f:
                            bad.append("休息要用 COUNTBLANK(按天数取范围) 或 =本月天数−其他")
                    if key in ("早班", "晚班", "年假"):
                        if "COUNTIF(" not in f or f'"{key[0]}"' not in f:
                            bad.append(f"没数「{key[0]}」")
                    if key == "实到":
                        refs_ok = stat_ref("早班") and stat_ref("早班") in f and stat_ref("晚班") and stat_ref("晚班") in f
                        quotes_ok = '"早"' in f and '"晚"' in f
                        if not (refs_ok or quotes_ok):
                            bad.append("实到没算早+晚")
                if bad:
                    report.add("六 数据统计", "error", f"{e['name']}「{key}」公式有问题（{'；'.join(bad)}）：{got}")
                else:
                    formula_ok += 1
            else:
                got_int = as_int(got)
                if key == "年假" and got_int is None:
                    pass  # 真表年假列是余额表达式（"2天5小时=5天-…"），只提醒
                elif got_int != want[key]:
                    report.add("六 数据统计", "error",
                               f"{e['name']} 统计「{key}」={got}，矩阵实际 {want[key]}")
                else:
                    value_ok += 1
    report.add("六 数据统计", "info",
               f"统计列：公式通过 {formula_ok} 处，数值比对 {value_ok} 处"
               + ("（公式按结构校验，不重算）" if formula_ok else ""))

    # 应到要全组统一（政策值）
    arrived = {e["name"]: e["stats"].get("应到") for e in cur["employees"] if e["stats"].get("应到") is not None}
    vals = {str(v).strip() for v in arrived.values() if not is_formula(v)}
    if vals and len(vals) > 1:
        report.add("六 数据统计", "error", f"应到不统一：{arrived}（应是本月天数 − 本月休息天数）")

    # 汇总行
    for group, members in groups_of(cur["employees"]).items():
        kinds = {"早班": "早", "晚班": "晚", "休息": ""}
        first, last = members[0]["row"], members[-1]["row"]
        for kind, code in kinds.items():
            row = cur["summary"].get((group, kind))
            if not row:
                report.add("六 数据统计", "warn", f"缺{group}{kind}汇总行")
                continue
            for d in days:
                want = sum(1 for e in members if e["shifts"][d] == code)
                got = row.get(d)
                letter = get_column_letter(2 + d)
                span = norm(f"{letter}{first}:{letter}{last}")
                if got is None:
                    report.add("六 数据统计", "error", f"{group}{kind}汇总 d{d} 为空，实际 {want}")
                elif is_formula(got):
                    f = norm(got)
                    bad = []
                    if span not in f:
                        bad.append(f"范围不是 {letter}{first}:{letter}{last}")
                    if kind == "休息":
                        if "COUNTBLANK(" not in f:
                            bad.append("休息没用 COUNTBLANK")
                    elif "COUNTIF(" not in f or f'"{code}"' not in f:
                        bad.append(f"没数「{code}」")
                    if bad:
                        report.add("六 数据统计", "error", f"{group}{kind}汇总 d{d} 公式有问题（{'；'.join(bad)}）")
                elif str(got).strip() != str(want):
                    report.add("六 数据统计", "error", f"{group}{kind}汇总 d{d}={got}，实际 {want}")

    # 上班人数
    if not cur["head_count"]:
        report.add("六 数据统计", "warn", "没找到「上班人数」行")
    else:
        for d in days:
            want = sum(1 for e in cur["employees"] if e["shifts"][d] in ("早", "晚"))
            got = cur["head_count"].get(d)
            letter = get_column_letter(2 + d)
            if got is None:
                report.add("六 数据统计", "error", f"上班人数 d{d} 为空，实际 {want}")
            elif is_formula(got):
                f = norm(got)
                ok = all(f"{letter}{r}" in f for r in (18, 19, 22, 23))
                if not ok:
                    report.add("六 数据统计", "error", f"上班人数 d{d} 公式没引用汇总行：{got}")
            elif str(got).strip() != str(want):
                report.add("六 数据统计", "error", f"上班人数 d{d}={got}，实际（早+晚）{want}")

    # 小计
    for group, cells in cur["subtotal"].items():
        members = [e for e in cur["employees"] if e["group"] == group]
        if not members:
            continue
        code_of = {"年假": "年", "休息": "", "早班": "早", "晚班": "晚"}
        first, last = members[0]["row"], members[-1]["row"]
        for key in ("年假", "休息", "早班", "晚班"):
            got = cells.get(key)
            if got is None:
                continue
            if is_formula(got):
                f = norm(got)
                letter = get_column_letter(35 + STAT_KEYS.index(key))
                if "SUM(" not in f or norm(f"{letter}{first}:{letter}{last}") not in f:
                    report.add("六 数据统计", "error", f"{group}小计「{key}」公式有问题：{got}")
            else:
                want = sum(1 for e in members for d in days if e["shifts"][d] == code_of[key])
                if str(got).strip() != str(want):
                    report.add("六 数据统计", "error", f"{group}小计「{key}」={got}，实际 {want}")


def check_carry(cur, prev, report):
    if not prev:
        report.add("六 数据统计", "info", "未给上月表，剩余假继承无法自动核对（人工看清单）")
        return
    pmap = {e["name"]: e for e in prev["employees"]}
    lines = []
    for e in cur["employees"]:
        pe = pmap.get(e["name"])
        now = text(e["stats"].get("剩余"))
        before = text(pe["stats"].get("剩余")) if pe else "?"
        lines.append(f"{e['name']} {before or '—'} → {now or '—'}")
    report.add("六 数据统计", "info", "剩余假继承（上月 → 本月，人工核对）：" + "；".join(lines))


# ---------------------------------------------------------------- 主流程

def main() -> None:
    parser = argparse.ArgumentParser(description="按验收清单校验排班表（只读）")
    parser.add_argument("--xlsx", required=True, help="本月排班表（xlsx 或 TSV）")
    parser.add_argument("--sheet", help="工作表名（默认活动表）")
    parser.add_argument("--prev", help="上月表（xlsx 或快照 TSV），用于跨月衔接校验")
    parser.add_argument("--lead", default="李守耀", help="售后组长（全早、休日全员上班）")
    parser.add_argument("--green", default=GREEN, help=f"值班绿标颜色，默认 {GREEN}")
    parser.add_argument("--max-streak", type=int, default=5, help="连续上班上限（默认 5）")
    parser.add_argument("--json", help="把问题清单写到 JSON")
    args = parser.parse_args()

    cur = parse(load_grid(args.xlsx, args.sheet))
    prev = parse(load_grid(args.prev)) if args.prev else None

    report = Report()
    if prev:
        check_cross(cur, prev, report, args.max_streak)
    else:
        report.add("〇 跨月衔接", "warn", "未给 --prev，跨月衔接未校验（务必人工看）")
    check_coverage(cur, report, args.lead)
    check_shifts(cur, report, args.max_streak)
    check_duty(cur, report, args.green.upper())
    check_after_duty(cur, report, args.lead)
    check_marks(cur, report)
    check_stats(cur, report)
    check_carry(cur, prev, report)

    symbol = {"error": "✗", "warn": "!", "info": "·"}
    section = None
    for sec, level, msg in report.items:
        if sec != section:
            print(f"\n== {sec} ==")
            section = sec
        print(f"  {symbol[level]} {msg}")

    errors = report.errors()
    warns = [i for i in report.items if i[1] == "warn"]
    print(f"\n验收结果：{'全部通过（自动项）' if not errors else f'{len(errors)} 项未通过'}"
          + (f"，{len(warns)} 项提醒" if warns else ""))
    print("提醒：剩余假口径、备注、请假/制度类条目请对照《验收清单.md》人工确认。")

    if args.json:
        with open(args.json, "w", encoding="utf-8") as handle:
            json.dump({"errors": [i[2] for i in errors],
                       "warns": [i[2] for i in warns],
                       "items": report.items}, handle, ensure_ascii=False, indent=2)
        print(f"已写出 JSON: {args.json}")


if __name__ == "__main__":
    main()
