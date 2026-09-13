# -*- coding: utf-8 -*-
"""单用途工具：把「排班计划 JSON」写回排班 xlsx（值 + 底色 + 公式统计 + 缓存值 + TSV）。

做成这样是为了能组合、能复用：
    解排班（人或脚本）→ 计划 JSON → 本工具写表 → 验排班表.py 验收

用法：
    cd "D:\\桌面\\办公软件\\21.AI自动排班经验"
    python "工具\\写排班表.py" --xlsx "测试数据\\2026年10月客服排班表.xlsx" `
        --plan "测试数据\\2026年10月-计划.json"
    # 可选：--tsv 输出TSV、--sheet 工作表名、--no-cache 不写公式缓存值

计划 JSON（`解排班示例.py` 会直接输出这个格式）：
{
  "meta": {
    "days": 31,                  # 当月天数
    "policy_rest": 6,            # 本月休息天数（政策值，写进 18 行右侧的「本月休息天数：」）
    "policy_note": "大小周",
    "lead": "李守耀",             # 售后组长：全早 + 黄色
    "carry": {"韩欢欢": "0", ...},           # 剩余：**上月结转**（数字=天数，或 "3天4小时"/"-3小时"）
    "annual": {"韩欢欢": 0, ...},             # 年假：上月结转的年假余额（同样支持 "1天3小时"；可省略）
    "codes": {"缪婷婷|8": "行"},                # 特殊格（只应落在休息格）
    "white": ["麦诺谦|15"],                     # 显式白底
    "carry_width": 4/省略                       # 可选：剩余列宽；默认跟每日排班的格宽（列 C）
  },
  "seller": {"韩欢欢": ["早","早","晚","",...]},  # 售前：长度 = days，空串 = 休
  "after":  {"李守耀": [...]},
  "duty":   {"韩欢欢|1": "早", ...},             # 售前值班（浅绿）：每天 早1 + 晚1
  "duty_after": {"缪婷婷|1": "晚", ...}          # 售后值班（浅蓝）：每班只 1 人；组长在岗时他自己是负责人（黄）
}

写完后表里：
    公式列：早班/晚班 = COUNTIF($C4:$AG4,...)；休息 = 「本月天数」−早−晚−年−行；
    实到 = 早+晚+行；应到 = 「本月天数」−「本月休息天数」。范围直接写 1号到31号（整月），
    天数靠政策格自动算——30 天的月份不会把 31 号空白算成休息。汇总行/小计/上班人数 也都是公式。
    结转列（值，不是公式）：年假 = 上月年假余额（原样写）；
    剩余 = 上月结转 + (本月休息天数 − 实休) × 8 小时（写成 "x天x小时"，跟公司工具一致）。
    **少休要补回来**：实休 5 天、应休 6 天 → 剩余 +1 天（不然就欠他的）；超休则记负数（写「（下个月）」，下月扣）。
    月份横幅（1 号~当月最后一天合并成 1 格写「N月」）+ 星期行大小周红标（字体颜色）。
    同时把公式算好的值缓存进 xlsx，不重算的程序（openpyxl 等）也能读到数字。

注：橙色（公休/调休）是**人工手标**的，我们排班不产生、也不涂——要标请手动标。
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Color, Font, PatternFill
from openpyxl.utils import get_column_letter
from copy import copy as _copy

STAT_KEYS = ("剩余", "年假", "休息", "早班", "晚班", "实到", "应到")
KINDS = ("早班", "晚班", "休息")

F_NONE = PatternFill(fill_type=None)
F_GREEN = PatternFill("solid", fgColor="E2F0D9")   # 售前值班
F_YELLOW = PatternFill("solid", fgColor="FFFF00")  # 售后值班组长（在岗）
F_BLUE = PatternFill("solid", fgColor="BDD7EE")    # 售后值班（每班只 1 人）
F_WHITE = PatternFill("solid", fgColor="FFFFFF")   # 显式白底
HOURS_PER_DAY = 8   # 假期按 8 小时/天折算（跟公司工具一致）


def parse_leave(v) -> float:
    """把「剩余/年假」列里的写法解析成小时：数字 = 天数；支持 "3天4小时"、"-3小时"、"0.5天"、"1.5小时"。
    真表里会写「余额=算式」，只取等号前的余额。"""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v) * HOURS_PER_DAY
    t = str(v).strip()
    if not t or t == "0":
        return 0.0
    # 真表里会写「余额=算式」（如 "-1天=3天3小时-1天-…（下个月）"）——只取等号前的余额
    t = t.split("=")[0].strip()
    if not t:
        return 0.0
    neg = t.startswith("-")
    if neg:
        t = t[1:]
    total = 0.0
    m = re.search(r"(\d+(?:\.\d+)?)\s*天", t)
    if m:
        total += float(m.group(1)) * HOURS_PER_DAY
    m = re.search(r"(\d+(?:\.\d+)?)\s*小时", t)
    if m:
        total += float(m.group(1))
    return -total if neg else total


def format_leave(hours: float) -> str:
    """小时 → "x天x小时" / "0" / "-x小时"（与公司工具的 格式化假期小时数 一致）。"""
    if abs(hours) < 1e-9:
        return "0"
    neg = hours < 0
    h = abs(hours)
    days = int(h // HOURS_PER_DAY)
    rem = h % HOURS_PER_DAY
    if rem == int(rem):
        rem = int(rem)
    if days == 0:
        out = f"{rem}小时"
    elif rem == 0:
        out = f"{days}天"
    else:
        out = f"{days}天{rem}小时"
    return f"-{out}" if neg else out


def text(v) -> str:
    return str(v).strip() if v is not None else ""


def split_keys(items) -> set[tuple[str, int]]:
    out = set()
    for item in items or []:
        name, _, day = str(item).partition("|")
        if day.isdigit():
            out.add((name.strip(), int(day)))
    return out


RED_FONT, BLACK_FONT = "FFFF0000", "FF000000"     # 大小周：红 = 休息日（字体颜色，不是底色）
F_BANNER = PatternFill("solid", fgColor="FF0000")  # 月份横幅（红底白字）


def month_label(ws, meta) -> str | None:
    """月份横幅文字：优先 meta.month，其次从标题（如 "2026年10月客服排班表"）里取。"""
    if meta.get("month"):
        return str(meta["month"])
    for r in range(1, 4):
        for c in range(1, 6):
            mo = re.search(r"(\d{1,2})月", text(ws.cell(r, c).value))
            if mo:
                return f"{int(mo.group(1))}月"
    return None


def find_weekday_row(ws, L) -> int | None:
    """星期行：日期行上/下一行，日列上写的是 一~日。"""
    cols = [L["day_col"][d] for d in L["day_col"]]
    for r in (L["header"] - 1, L["header"] + 1):
        if r < 1:
            continue
        vals = {text(ws.cell(r, c).value) for c in cols} - {""}
        if vals and vals <= set("一二三四五六日"):
            return r
    return None


def write_month_banner(ws, row: int, col0: int, col1: int, label: str) -> None:
    """把 1 号~当月最后一天 的日列在横幅行上合成 1 格（醒目，避免客服看错月份）。"""
    for rng in list(ws.merged_cells.ranges):      # 先拆掉与日列相交的旧合并（真表可能是一排散格）
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col1 and rng.max_col >= col0:
            ws.unmerge_cells(str(rng))
    for c in range(col0, col1 + 1):
        cell = ws.cell(row, c)
        cell.value = None
        cell.fill = F_BANNER
        cell.font = Font(bold=True, color="FFFFFFFF", size=10)
    ws.merge_cells(start_row=row, start_column=col0, end_row=row, end_column=col1)
    top = ws.cell(row, col0)
    top.value = label
    top.alignment = Alignment(horizontal="center", vertical="center")


def mark_weekend_reds(ws, row: int, day_col: dict[int, int], days: list[int], red_days: set[int]) -> None:
    """星期行按大小周标红（字体颜色）：周日永远红，周六按大周红/小周黑，工作日黑。"""
    for d in days:
        c = day_col.get(d)
        if c:
            cell = ws.cell(row, c)
            font = _copy(cell.font)                 # 只改字体颜色，其余（字体/字号/加粗）保持原样
            font.color = Color(rgb=RED_FONT if d in red_days else BLACK_FONT)
            cell.font = font


def find_layout(ws, plan) -> dict:
    """按标准模板找位置：日期行 → 天列；姓名 → 行；汇总/上班人数/小计行。"""
    header = name_col = None
    day_col: dict[int, int] = {}
    for r in range(1, min(ws.max_row, 20) + 1):
        for c in range(1, ws.max_column + 1):
            if text(ws.cell(r, c).value) == "日期":
                header, name_col = r, c
                break
        if header:
            break
    if not header:
        raise SystemExit("没找到含「日期」的表头行")
    for c in range(1, ws.max_column + 1):
        v = text(ws.cell(header, c).value)
        if v.isdigit() and 1 <= int(v) <= 31:
            day_col.setdefault(int(v), c)

    stat_col = {}
    for r in range(1, min(ws.max_row, 8) + 1):
        for c in range(name_col + 1, ws.max_column + 1):
            v = text(ws.cell(r, c).value)
            if v in STAT_KEYS:
                stat_col.setdefault(v, c)
    if "早班" not in stat_col or "休息" not in stat_col:
        raise SystemExit("没找到统计表头（剩余/年假/休息/早班/晚班/实到/应到）")

    def row_of(name: str) -> int:
        for r in range(header + 1, ws.max_row + 1):
            if text(ws.cell(r, name_col).value) == name:
                return r
        raise SystemExit(f"表里找不到「{name}」这一行")

    seller_rows = {p: row_of(p) for p in plan.get("seller", {})}
    after_rows = {p: row_of(p) for p in plan.get("after", {})}
    summary, group = {}, ""
    for r in range(header + 1, ws.max_row + 1):
        g1, g2 = text(ws.cell(r, 1).value), text(ws.cell(r, 2).value)
        if g1 in ("售前", "售后") and g2 in KINDS:
            group, kind = g1, g2
        elif g2 in KINDS:
            kind = g2
        else:
            kind = ""
        if kind and group:
            summary[(group, kind)] = r
    head_row = next((r for r in range(header + 1, ws.max_row + 1)
                     if text(ws.cell(r, 2).value) == "上班人数"), None)
    if head_row is None:
        head_row = summary[("售后", "休息")] + 2
    subtotal_row = max(seller_rows.values()) + 1
    return {"header": header, "name_col": name_col, "day_col": day_col, "stat_col": stat_col,
            "seller": seller_rows, "after": after_rows, "summary": summary,
            "head_row": head_row, "subtotal_row": subtotal_row}


def main() -> None:
    ap = argparse.ArgumentParser(description="把计划 JSON 写回排班表（值+底色+公式+缓存值+TSV）")
    ap.add_argument("--xlsx", required=True, help="排班表 xlsx（就地修改，建议先备份）")
    ap.add_argument("--plan", required=True, help="计划 JSON")
    ap.add_argument("--sheet", help="工作表名（默认活动表）")
    ap.add_argument("--tsv", help="可选：同时导出 TSV（数值，不带公式）")
    ap.add_argument("--no-cache", action="store_true", help="不写公式缓存值")
    args = ap.parse_args()

    plan = json.loads(Path(args.plan).read_text(encoding="utf-8"))
    meta = plan.get("meta", {})
    seller, after = plan.get("seller", {}), plan.get("after", {})
    days_all = sorted(plan.get("meta", {}).get("days_list") or range(1, (meta.get("days") or 31) + 1))
    lead = meta.get("lead", "")
    codes = {tuple(k.split("|")[:1]) + (int(k.split("|")[1]),): v for k, v in (meta.get("codes") or {}).items()}
    white = split_keys(meta.get("white"))
    duty = {tuple(k.split("|")[:1]) + (int(k.split("|")[1]),) for k in (plan.get("duty") or {})}
    duty_after = {tuple(k.split("|")[:1]) + (int(k.split("|")[1]),) for k in (plan.get("duty_after") or {})}
    carry = meta.get("carry") or {}

    wb = load_workbook(args.xlsx)
    ws = wb[args.sheet] if args.sheet else wb.active
    L = find_layout(ws, plan)
    days = [d for d in days_all if d in L["day_col"]]

    values: dict[tuple[int, int], object] = {}   # 公式缓存值 / TSV 值

    # ---- 表头：月份横幅（合并）+ 星期行大小周红标 ----
    banner_row = (min(L["after"].values()) - 1) if L["after"] else None
    busy = set(L["seller"].values()) | set(L["after"].values()) | {L["header"], L["subtotal_row"], L["head_row"]}
    if banner_row and banner_row > L["header"] and banner_row not in busy:
        label = month_label(ws, meta)
        if label:
            write_month_banner(ws, banner_row, L["day_col"][days_all[0]], L["day_col"][days[-1]], label)
            print(f"月份横幅：第 {banner_row} 行 {get_column_letter(L['day_col'][days_all[0]])}~{get_column_letter(L['day_col'][days[-1]])} 合并为「{label}」")
    wd_row = find_weekday_row(ws, L)
    red_days = {int(d) for d in (meta.get("weekend_red_days") or [])}
    if wd_row and red_days:
        mark_weekend_reds(ws, wd_row, L["day_col"], days, red_days)
        print(f"星期行（第 {wd_row} 行）按大小周标红 {len(red_days & set(days))} 天")
    elif wd_row:
        print("提示：计划里没写 weekend_red_days，星期行大小周红标没动（可用 工具/推应休天数.py 推出）")

    def write_value(r, c, v):
        ws.cell(r, c).value = v
        values[(r, c)] = v

    # 天数格（政策格）先算出来，统计公式要引用
    rest_row = L["summary"][("售前", "早班")]
    days_row = L["summary"][("售前", "晚班")]
    label_col = L["stat_col"]["早班"]           # 真表把政策格放在统计区右侧
    val_col = label_col + 1
    day_cell = f"${get_column_letter(val_col)}${days_row}"     # 「本月天数：」的值
    rest_cell = f"${get_column_letter(val_col)}${rest_row}"    # 「本月休息天数：」的值
    n_days = len(days)

    # ---- 班次格（值 + 底色）----
    matrix: dict[tuple[str, int], object] = {}
    per_person: dict[str, tuple[int, int, int, int, int]] = {}
    for group, shifts in (("seller", seller), ("after", after)):
        for p, row in ((group == "seller" and seller) or after).items():
            r = L["seller" if group == "seller" else "after"][p]
            nian = xing = early = late = rest = 0
            for i, s in enumerate(row, start=1):
                d = days_all[i - 1] if i - 1 < len(days_all) else i
                c = L["day_col"].get(d)
                if not c:
                    continue
                code = codes.get((p, d))
                val = code if code else (s or None)
                cell = ws.cell(r, c)
                cell.value = val
                values[(r, c)] = val
                if group == "seller":
                    cell.fill = F_GREEN if (p, d) in duty else (F_WHITE if (p, d) in white else F_NONE)
                elif not s:                                # 售后休息：不涂色（橙色是人工手标的，我们排班不产生）
                    cell.fill = F_NONE
                elif p == lead:                            # 组长在岗 = 值班负责人（黄）
                    cell.fill = F_YELLOW
                elif (p, d) in duty_after:                 # 售后值班：每班只 1 人（浅蓝）
                    cell.fill = F_BLUE
                elif (p, d) in white:
                    cell.fill = F_WHITE
                else:
                    cell.fill = F_NONE
                if code == "年":
                    nian += 1
                elif code == "行":
                    xing += 1
                elif s == "早":
                    early += 1
                elif s == "晚":
                    late += 1
                else:
                    rest += 1
                matrix[(p, d)] = val
            per_person[p] = (nian, rest, early, late, early + late + xing)

    # ---- 统计列：公式（早/晚/休息/实到/应到）+ 结转列（年假/剩余，值）----
    annual = meta.get("annual") or {}
    policy_rest = int(meta.get("policy_rest") or 0)
    for p, (nian, rest, early, late, arrived) in per_person.items():
        r = L["seller"].get(p) or L["after"][p]
        stat = L["stat_col"]
        col = {k: get_column_letter(stat[k]) for k in ("休息", "早班", "晚班")}
        rng = f"$C{r}:$AG{r}"          # 整月范围：1号到31号（30 天的月份 31 号列是空的，不参与计数）
        if p in carry:
            # 剩余 = 上月结转 + (本月应休 − 实际排休) × 8 小时（公司口径，step_28_update_remaining）
            # 超休（负数）就记作休下个月的假，备注「（下个月）」（跟真表一样）
            left = parse_leave(carry[p]) + (policy_rest - rest) * HOURS_PER_DAY
            text = format_leave(left) + ("（下个月）" if left < 0 else "")
            ws.cell(r, stat["剩余"]).value = text
            values[(r, stat["剩余"])] = text
        if p in annual:
            # 年假 = 上月结转的年假余额（原样写，不是本月计数；本月真休了年假要人工扣减）
            ws.cell(r, stat["年假"]).value = annual[p]
            values[(r, stat["年假"])] = annual[p]
        formulas = {
            "早班": f'=COUNTIF({rng},"早")',
            "晚班": f'=COUNTIF({rng},"晚")',
            "休息": f'={day_cell}-{col["早班"]}{r}-{col["晚班"]}{r}-COUNTIF({rng},"年")-COUNTIF({rng},"行")',
            "实到": f'={col["早班"]}{r}+{col["晚班"]}{r}+COUNTIF({rng},"行")',
            "应到": f"={day_cell}-{rest_cell}",
        }
        for key, f in formulas.items():
            ws.cell(r, stat[key]).value = f
            values[(r, stat[key])] = {"休息": rest, "早班": early, "晚班": late, "实到": arrived,
                                      "应到": n_days - policy_rest}[key]

    # ---- 售前小计：SUM（年假/剩余 是结转余额文本，不做小计）----
    srows = sorted(L["seller"].values())
    for key in ("休息", "早班", "晚班"):
        col = L["stat_col"][key]
        letter = get_column_letter(col)
        ws.cell(L["subtotal_row"], col).value = f"=SUM({letter}{srows[0]}:{letter}{srows[-1]})"
        values[(L["subtotal_row"], col)] = sum(
            per_person[p][{"休息": 1, "早班": 2, "晚班": 3}[key]] for p in seller)
    for key in ("剩余", "年假", "实到", "应到"):
        c = L["stat_col"][key]
        cell = ws.cell(L["subtotal_row"], c)
        if key in ("实到", "应到"):
            cell.value = None
            values[(L["subtotal_row"], c)] = None

    # ---- 汇总行 + 上班人数：公式 ----
    sums = {}
    for d in days:
        c = L["day_col"][d]
        letter = get_column_letter(c)
        se = sum(1 for p in seller if matrix.get((p, d)) == "早")
        sl = sum(1 for p in seller if matrix.get((p, d)) == "晚")
        sf = sum(1 for p in seller if matrix.get((p, d)) is None)
        ae = sum(1 for p in after if matrix.get((p, d)) == "早")
        al = sum(1 for p in after if matrix.get((p, d)) == "晚")
        af = sum(1 for p in after if matrix.get((p, d)) is None)
        work = se + sl + ae + al
        for kind, f, v in (
            (("售前", "早班"), f'=COUNTIF({letter}${srows[0]}:{letter}${srows[-1]},"早")', se),
            (("售前", "晚班"), f'=COUNTIF({letter}${srows[0]}:{letter}${srows[-1]},"晚")', sl),
            (("售前", "休息"), f"=COUNTBLANK({letter}${srows[0]}:{letter}${srows[-1]})", sf),
            (("售后", "早班"), f'=COUNTIF({letter}${sorted(L["after"].values())[0]}:{letter}${sorted(L["after"].values())[-1]},"早")', ae),
            (("售后", "晚班"), f'=COUNTIF({letter}${sorted(L["after"].values())[0]}:{letter}${sorted(L["after"].values())[-1]},"晚")', al),
            (("售后", "休息"), f"=COUNTBLANK({letter}${sorted(L['after'].values())[0]}:{letter}${sorted(L['after'].values())[-1]})", af),
        ):
            row = L["summary"][kind]
            ws.cell(row, c).value = f
            values[(row, c)] = v
        ws.cell(L["head_row"], c).value = (f"={letter}{L['summary'][('售前','早班')]}+{letter}{L['summary'][('售前','晚班')]}"
                                           f"+{letter}{L['summary'][('售后','早班')]}+{letter}{L['summary'][('售后','晚班')]}")
        values[(L["head_row"], c)] = work
        sums[d] = (se, sl, sf, ae, al, af, work)

    # ---- 政策格 + 标签 ----
    ws.cell(L["head_row"], 2).value = "上班人数"
    values[(L["head_row"], 2)] = "上班人数"
    labels = {
        rest_row: ("本月休息天数：", int(meta.get("policy_rest") or 0), meta.get("policy_note", "")),
        days_row: ("本月天数：", None, ""),
    }
    for r, (label, v, note) in labels.items():
        ws.cell(r, label_col).value = label
        values[(r, label_col)] = label
        if label.startswith("本月天数"):
            ws.cell(r, val_col).value = "=COUNTA($C$2:$AG$2)"
            values[(r, val_col)] = n_days
        else:
            ws.cell(r, val_col).value = v
            values[(r, val_col)] = v
        if note:
            ws.cell(r, val_col + 1).value = note
            values[(r, val_col + 1)] = note

    # ---- 列宽（剩余列跟着每日排班的格宽，别自己乱定）----
    carry_col = L["stat_col"]["剩余"]
    day_letter = get_column_letter(min(L["day_col"].values()))
    day_width = ws.column_dimensions[day_letter].width or 4
    ws.column_dimensions[get_column_letter(carry_col)].width = float(meta.get("carry_width") or day_width)
    for key in ("年假", "休息", "早班", "晚班", "实到", "应到"):
        ws.column_dimensions[get_column_letter(L["stat_col"][key])].width = max(
            9.0, ws.column_dimensions[get_column_letter(L["stat_col"][key])].width or 0)

    wb.calculation.fullCalcOnLoad = True   # 用真表格打开时重算一遍
    wb.save(args.xlsx)
    print(f"xlsx saved: {args.xlsx}")

    if not args.no_cache:
        n = inject_cached_values(args.xlsx, ws.title, {k: v for k, v in values.items()
                                                       if isinstance(v, (int, float)) and v is not None})
        print(f"公式缓存值已写入 {n} 个单元格（不重算的读法也能读到数字）")

    if args.tsv:
        lines = []
        for r in range(1, ws.max_row + 1):
            cells = []
            for c in range(1, max(42, ws.max_column) + 1):
                v = values.get((r, c), ws.cell(r, c).value)
                if isinstance(v, str) and v.startswith("="):
                    v = None
                cells.append("" if v is None else str(v))
            lines.append("\t".join(cells))
        with open(args.tsv, "w", encoding="utf-8-sig", newline="") as fh:
            fh.write("\r\n".join(lines) + "\r\n")
        print(f"tsv saved: {args.tsv}")

    print("\n每日覆盖(售前早/晚/休 | 售后早/晚/休 | 上班人数):")
    for d in days:
        print(f"  d{d:2d}: {sums[d][0]}/{sums[d][1]}/{sums[d][2]} | {sums[d][3]}/{sums[d][4]}/{sums[d][5]} | {sums[d][6]}")
    print("\n每人统计（公式应得值）:")
    for p, (nian, rest, early, late, arrived) in per_person.items():
        print(f"  {p}: 年{nian} 休{rest} 早{early} 晚{late} 实到{arrived} 应到{n_days - int(meta.get('policy_rest') or 0)}")


def inject_cached_values(path: str, sheet_name: str, cells: dict[tuple[int, int], float]) -> int:
    """openpyxl 写公式只留空 <v></v>；这里把算好的值填进去当缓存，供不重算的读法使用。"""
    if not cells:
        return 0
    path = str(path)
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        wb_xml = zf.read("xl/workbook.xml").decode("utf-8")
        rels_xml = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        hit = re.search(r'<sheet[^>]*name="%s"[^>]*r:id="([^"]+)"' % re.escape(sheet_name), wb_xml)
        if not hit:
            return 0
        rid = hit.group(1)
        target = None
        for tag in re.findall(r"<Relationship\b[^>]*/>", rels_xml):
            if f'Id="{rid}"' in tag:
                found = re.search(r'Target="([^"]+)"', tag)
                target = found.group(1) if found else None
                break
        if not target:
            return 0
        target = target.lstrip("/")
        sheet_path = target if target.startswith("xl/") else "xl/" + target
        sheet_xml = zf.read(sheet_path).decode("utf-8")
        count = 0
        for (r, c), value in cells.items():
            ref = f"{get_column_letter(c)}{r}"
            pat = re.compile(r'(<c r="%s"[^>]*>(?:<f>.*?</f>))<v></v>' % ref)
            num = int(value) if float(value).is_integer() else value
            sheet_xml, n = pat.subn(lambda m: m.group(1) + f"<v>{num}</v>", sheet_xml, count=1)
            count += n
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
            for name in names:
                out.writestr(name, sheet_xml.encode("utf-8") if name == sheet_path else zf.read(name))
    shutil.move(tmp, path)
    return count


if __name__ == "__main__":
    main()
