# -*- coding: utf-8 -*-
"""推本月应休天数（大小周部分）：从上月表「星期行周六是否红」推出来，法定节假日仍要问人事/用户。

口径来自公司原版工具（`schedule_tools/steps/step_03_mark_weekends.py`：
"根据上月的周六标记推导本月的大小周，设置周六红色" + `excel_header_core.py`）：
    · 红 = **字体颜色**（不是底色）：周日永远红；周六红 = 大周（双休），周六黑 = 小周（单休）
    · 大小周必须**严格交替**（公司工具不交替会报错）
    · 推导：从源月最后一个周一开始，逐周翻转，直到目标月
    · 本月应休 = 大小周红标天数 + 法定节假日天数

用法：
    python "工具\\推应休天数.py" --prev "上月表.xlsx"                     # 自动推下一个月
    python "工具\\推应休天数.py" --prev "上月表.xlsx" --month 2026-10
    python "工具\\推应休天数.py" --prev "上月表.xlsx" --holidays 3         # 加上法定节假日（问过人事的值）
    python "工具\\推应休天数.py" --prev "上月表.xlsx" --phase 大           # 上月表没标红时，手工给上月最后一周类型

输出里的 `weekend_red_days` 可直接填进计划 JSON（`工具\\写排班表.py` 会照它给星期行标红）。

注意：**每月都要问用户/人事确认应休天数**（法定节假日 + 大小周，月月不同），不要沿用上月值。
"""
from __future__ import annotations

import argparse
import calendar
import re
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

RED_RGB = {"FFFF0000", "00FF0000", "FF0000"}
WEEKDAY_CHARS = "一二三四五六日"


def text(v) -> str:
    return str(v).strip() if v is not None else ""


def find_header(ws) -> tuple[int, int]:
    """返回（日期行, 起始列）：起始列 = 写「1」的那一列。"""
    for r in range(1, min(ws.max_row, 20) + 1):
        for c in range(1, ws.max_column + 1):
            if text(ws.cell(r, c).value) == "日期":
                return r, c + 1
    raise SystemExit("没找到含「日期」的表头行")


def find_weekday_row(ws, date_row: int, col0: int, days: int) -> int:
    for r in (date_row - 1, date_row + 1):
        if r < 1:
            continue
        vals = {text(ws.cell(r, col0 + i).value) for i in range(days)} - {""}
        if vals and vals <= set(WEEKDAY_CHARS):
            return r
    raise SystemExit("没找到星期行（日期行上/下一行的日列应写 一~日）")


def month_of_title(ws) -> tuple[int, int] | None:
    for r in range(1, 4):
        for c in range(1, 6):
            mo = re.search(r"(\d{4})\s*年\s*(\d{1,2})\s*月", text(ws.cell(r, c).value))
            if mo:
                return int(mo.group(1)), int(mo.group(2))
    return None


def is_red(cell) -> bool:
    color = cell.font.color if cell.font else None
    if not color or not color.rgb:
        return False
    return str(color.rgb).upper()[-8:] in RED_RGB or str(color.rgb).upper() in RED_RGB


def extract_big_week(ws, year: int, month: int) -> dict[date, bool]:
    """上月：每个含周六的周（键=周一）→ 周六是否红（True=大周/双休）。必须严格交替。"""
    date_row, col0 = find_header(ws)
    days = calendar.monthrange(year, month)[1]
    wd_row = find_weekday_row(ws, date_row, col0, days)
    week_type: dict[date, bool] = {}
    for d in range(1, days + 1):
        cur = date(year, month, d)
        if cur.weekday() != 5:                      # 只看周六
            continue
        week_type[cur - timedelta(days=cur.weekday())] = is_red(ws.cell(wd_row, col0 + d - 1))
    if not week_type:
        raise SystemExit(f"{year}年{month}月没找到任何周六")
    weeks = sorted(week_type)
    for a, b in zip(weeks, weeks[1:]):
        if week_type[a] == week_type[b]:
            raise SystemExit(f"大小周没有交替（{a} 与 {b} 的周六同为{'红' if week_type[a] else '黑'}）——"
                             "请检查上月表星期行的红色标记")
    return week_type


def derive(year: int, month: int, src: dict[date, bool]) -> dict[date, bool]:
    """由源月外推到目标月（公司 excel_header_core.推导目标月周类型 同款算法）。"""
    first, last = date(year, month, 1), date(year, month, calendar.monthrange(year, month)[1])
    cursor = max(src)
    big = src[cursor]
    first_monday = first - timedelta(days=first.weekday())
    while cursor < first_monday:
        cursor += timedelta(days=7)
        big = not big
    last_monday = last - timedelta(days=last.weekday())
    out: dict[date, bool] = {}
    while cursor <= last_monday:
        out[cursor] = big
        cursor += timedelta(days=7)
        big = not big
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description="从上月表的周六红标推本月大小周休息天数（法定节假日仍要问人事/用户）")
    ap.add_argument("--prev", required=True, help="上月排班表 xlsx（要有字体颜色；TSV 读不到）")
    ap.add_argument("--sheet", help="工作表名（默认活动表）")
    ap.add_argument("--month", help="目标月份 YYYY-MM（默认 = 上月+1）")
    ap.add_argument("--holidays", type=int, default=None, help="法定节假日天数（问过人事/用户的值）")
    ap.add_argument("--phase", choices=["大", "小"], help="上月表没标红时，手工给上月最后一个含周六周的类型")
    args = ap.parse_args()

    wb = load_workbook(args.prev)                  # 需要样式，不能用 data_only
    ws = wb[args.sheet] if args.sheet else wb.active

    title = month_of_title(ws)
    if not title:
        raise SystemExit("标题里没找到「YYYY年M月」（如 2026年9月客服排班表），请确认这是排班表")
    src_year, src_month = title

    if args.phase:                                 # 手工给相位：构造一个只含最后一周的源
        days = calendar.monthrange(src_year, src_month)[1]
        last_day = date(src_year, src_month, days)
        sat = last_day - timedelta(days=(last_day.weekday() - 5) % 7)
        src = {sat - timedelta(days=sat.weekday()): args.phase == "大"}
        print(f"（手工相位）上月最后一周按{'大周（双休）' if args.phase == '大' else '小周（单休）'}算")
    else:
        src = extract_big_week(ws, src_year, src_month)

    if args.month:
        mo = re.fullmatch(r"(\d{4})-(\d{1,2})", args.month)
        if not mo:
            raise SystemExit("--month 要写成 YYYY-MM，如 2026-10")
        year, month = int(mo.group(1)), int(mo.group(2))
    else:
        year, month = (src_year + 1, 1) if src_month == 12 else (src_year, src_month + 1)

    print(f"上月 {src_year}年{src_month}月（周六红标）：")
    for monday in sorted(src):
        sat = monday + timedelta(days=5)
        print(f"    {monday}~{monday + timedelta(days=6)}  {sat.month}月{sat.day}日(周六) "
              f"{'红 → 大周（双休）' if src[monday] else '黑 → 小周（单休）'}")

    weeks = derive(year, month, src)
    days = calendar.monthrange(year, month)[1]
    red_days: list[int] = []
    print(f"\n本月 {year}年{month}月（按交替推出）：")
    for monday in sorted(weeks):
        big = weeks[monday]
        sun = monday + timedelta(days=6)
        sat = monday + timedelta(days=5)
        in_month = [d for d in (sat, sun) if d.month == month and d.year == year]
        marks = []
        for d in (monday + timedelta(days=i) for i in range(7)):
            if d.month == month and d.year == year:
                if d.weekday() == 6 or (d.weekday() == 5 and big):
                    red_days.append(d.day)
                    marks.append(f"{d.day}{WEEKDAY_CHARS[d.weekday()]}")
        span = f"{monday}~{monday + timedelta(days=6)}"
        print(f"    {span}  {'大周（双休）' if big else '小周（单休）'}  本月标红：{'、'.join(marks) if marks else '（不在本月）'}"
              f"{'' if in_month else '  ← 跨月周，只算本月内的天'}")

    red_days = sorted(set(red_days))
    n_big_week = len(red_days)
    print(f"\n大小周休息 = {n_big_week} 天（红标：{'、'.join(map(str, red_days))}）")
    print(f'计划 JSON 可直接用：  "weekend_red_days": {red_days},')
    if args.holidays is None:
        print("\n⚠️ 本月应休 = 大小周 %d 天 + 法定节假日 ? 天 —— 法定节假日部分**要问用户/人事**"
              "（月月不同，不要沿用上月）；确认后带 --holidays 再跑一次。" % n_big_week)
    else:
        print(f"\n本月应休 = 大小周 {n_big_week} 天 + 法定节假日 {args.holidays} 天 = {n_big_week + args.holidays} 天"
              f"（政策格 policy_rest 填 {n_big_week + args.holidays}）")
    print("⚠️ 每个月都要重新确认：法定节假日 + 大小周，月月不同。")


if __name__ == "__main__":
    main()
